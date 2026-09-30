"""국뽕 소재 찾기 + 대본 시트 테스트: python -m unittest tests.test_kukppong -v"""
import os
import tempfile
import unittest
from unittest import mock

os.environ["SHORTS_DB"] = os.path.join(tempfile.mkdtemp(), "test.db")
os.environ["YT_API_KEY"] = "fake-key"

from openpyxl import load_workbook  # noqa: E402

from shorts import config, kukppong, script_sheet, youtube  # noqa: E402
from shorts.db import connect  # noqa: E402
from tests import fake_youtube  # noqa: E402
from tests.fake_youtube import fake_get  # noqa: E402

# 가짜 데이터: v1 200만(구독 100만, x2) · v2 50만(구독 5천, x100) · v3 15만(구독 30만) · v4 10분 롱폼 900만 · v5 5만


@mock.patch("shorts.youtube.requests.get", side_effect=fake_get)
class KukppongTest(unittest.TestCase):
    def setUp(self):
        with connect() as con:
            for t in ("api_cache", "quota_log", "kukppong_videos"):
                con.execute(f"DELETE FROM {t}")

    def test_bench_is_shorts_sorted_by_outlier(self, _):
        rows = kukppong.collect("bench", min_views=100_000)
        self.assertEqual([r["video_id"] for r in rows], ["v2", "v1", "v3"])  # 롱폼 v4·5만회 v5 제외
        self.assertTrue(rows[0]["channel_url"].endswith("/channel/c2"))

    def test_source_includes_long_and_new_channels(self, _):
        rows = kukppong.collect("source", min_views=100_000, sort="views")
        self.assertEqual(rows[0]["video_id"], "v4")                   # 롱폼 포함
        new = kukppong.collect("source", min_views=0, max_subs=10_000)
        self.assertEqual({r["video_id"] for r in new}, {"v2", "v5"})   # 구독자 1만 이하만
        with connect() as con:
            self.assertEqual(con.execute("SELECT COUNT(*) FROM kukppong_videos").fetchone()[0], 5)

    def test_estimate(self, _):
        self.assertEqual(kukppong.estimate_units(["bench", "source"]), 8 * 102)
        self.assertEqual(kukppong.estimate_units(["source"], ["a"]), 102)


@mock.patch("shorts.youtube.requests.get", side_effect=fake_get)
class ScriptSheetTest(unittest.TestCase):
    def test_parse_video_id(self, _):
        p = script_sheet.parse_video_id
        self.assertEqual(p("https://youtube.com/shorts/D-UDWeCDv4Y"), "D-UDWeCDv4Y")
        self.assertEqual(p("https://www.youtube.com/watch?v=jfVVXYTZykw&t=10s"), "jfVVXYTZykw")
        self.assertEqual(p("https://youtu.be/jfVVXYTZykw"), "jfVVXYTZykw")
        self.assertEqual(p("jfVVXYTZykw"), "jfVVXYTZykw")
        self.assertIsNone(p("not a video"))

    def test_make_sheet(self, _):
        out = tempfile.mkdtemp()
        real_id = "AbCdEfGhI-_"  # 실제 유튜브 ID처럼 11자
        with mock.patch.object(script_sheet, "OUTPUT_DIR", config.Path(out)), \
                mock.patch.dict("tests.fake_youtube.VIDEOS", {real_id: fake_youtube.VIDEOS["v2"]}):
            made, missing = script_sheet.make_sheets([f"https://youtube.com/shorts/{real_id}", "zzzzzzzzzzz", "bad"])
        self.assertEqual(len(made), 1)
        self.assertEqual(missing, ["zzzzzzzzzzz", "bad"])
        wb = load_workbook(made[0])
        self.assertEqual(wb.sheetnames, ["대본", "작성 가이드"])
        ws = wb["대본"]
        self.assertEqual(ws["B2"].value, f"https://youtube.com/shorts/{real_id}")
        self.assertEqual(ws["B5"].value, "출처 - チャンネルc2")
        self.assertEqual(ws["A9"].value, "수정 전")
        self.assertEqual(ws["G9"].value, "수정 후")
        self.assertEqual(ws["B10"].value, "猫が逆さまに歩く")                       # 수정 전 제목 = 원본 제목
        self.assertEqual([ws.cell(12, c).value for c in range(1, 6)], script_sheet.COLS)
        self.assertEqual(ws["A13"].value, "기 (0~3초 후킹)")
        self.assertEqual(ws["G15"].value, "승 (몰입·갈등)")
        self.assertEqual(youtube.used_today() >= 2, True)


if __name__ == "__main__":
    unittest.main()
